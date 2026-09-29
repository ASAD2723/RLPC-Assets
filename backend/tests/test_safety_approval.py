"""Tests for the new Safety Team approval feature and RBAC around it."""
import io
import os
import pytest
import requests
import openpyxl

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://emp-purchases.preview.emergentagent.com").rstrip("/")
API = f"{BASE_URL}/api"

CREDS = {
    "biz_158": ("158", "Rlpc_974"),
    "biz_admin": ("IT admin", "Rlpc_974"),
    "safety_170": ("170", "Safety_91"),
    "safety_121": ("121", "Safety_974"),
    "editor_16": ("16", "Roc_974"),
    "editor_76": ("76", "Pro_974"),
    "editor_122": ("122", "Proc_974"),
    "editor_126": ("126", "Purchase_974"),
}


def _login(u, p):
    r = requests.post(f"{API}/auth/login", json={"username": u, "password": p}, timeout=30)
    assert r.status_code == 200, f"login {u} -> {r.status_code}: {r.text}"
    return r.json()


@pytest.fixture(scope="module")
def tokens():
    return {k: _login(u, p)["token"] for k, (u, p) in CREDS.items()}


def _h(tok):
    return {"Authorization": f"Bearer {tok}"}


# ---------- Login role flags ----------
class TestLoginFlags:
    def test_safety_170_flags(self):
        d = _login(*CREDS["safety_170"])
        assert d["can_approve"] is False and d["can_safety_approve"] is True

    def test_safety_121_flags(self):
        d = _login(*CREDS["safety_121"])
        assert d["can_approve"] is False and d["can_safety_approve"] is True

    def test_biz_158_flags(self):
        d = _login(*CREDS["biz_158"])
        assert d["can_approve"] is True and d["can_safety_approve"] is True

    def test_biz_admin_flags(self):
        d = _login(*CREDS["biz_admin"])
        assert d["can_approve"] is True and d["can_safety_approve"] is True

    @pytest.mark.parametrize("key", ["editor_16", "editor_76", "editor_122", "editor_126"])
    def test_editor_flags(self, key):
        d = _login(*CREDS[key])
        assert d["can_approve"] is False and d["can_safety_approve"] is False

    def test_me_endpoint(self, tokens):
        r = requests.get(f"{API}/auth/me", headers=_h(tokens["safety_170"]))
        assert r.status_code == 200
        j = r.json()
        assert j["username"] == "170" and j["can_safety_approve"] is True and j["can_approve"] is False

    def test_me_unauthenticated(self):
        assert requests.get(f"{API}/auth/me").status_code == 401


# ---------- Create purchase (helper) as biz for base records ----------
@pytest.fixture
def base_purchase(tokens):
    payload = {
        "employee_id": "TEST_S1", "employee_name": "TEST Safety",
        "purchase_type": "Mobile Purchase", "payment_mode": "Cash",
        "payment_by": "Jogy Joseph",
        "business_manager_approved": False, "safety_team_approved": False,
    }
    r = requests.post(f"{API}/purchases", json=payload, headers=_h(tokens["biz_158"]))
    assert r.status_code == 201, r.text
    pid = r.json()["id"]
    yield pid
    # cleanup
    requests.delete(f"{API}/purchases/{pid}", headers=_h(tokens["biz_158"]))


# ---------- PATCH /safety-approval RBAC ----------
class TestSafetyApprovalPatch:
    def test_no_token_401(self, base_purchase):
        r = requests.patch(f"{API}/purchases/{base_purchase}/safety-approval", json={"approved": True})
        assert r.status_code == 401

    def test_editor_forbidden(self, tokens, base_purchase):
        r = requests.patch(f"{API}/purchases/{base_purchase}/safety-approval",
                           json={"approved": True}, headers=_h(tokens["editor_16"]))
        assert r.status_code == 403

    @pytest.mark.parametrize("actor", ["safety_170", "safety_121", "biz_158", "biz_admin"])
    def test_allowed_users(self, tokens, base_purchase, actor):
        r = requests.patch(f"{API}/purchases/{base_purchase}/safety-approval",
                           json={"approved": True}, headers=_h(tokens[actor]))
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["safety_team_approved"] is True
        # business untouched
        assert body["business_manager_approved"] is False
        # GET verify persistence
        g = requests.get(f"{API}/purchases", params={"search": "TEST_S1"})
        items = [i for i in g.json()["items"] if i["id"] == base_purchase]
        assert items and items[0]["safety_team_approved"] is True
        assert items[0]["business_manager_approved"] is False
        # toggle back off
        r2 = requests.patch(f"{API}/purchases/{base_purchase}/safety-approval",
                            json={"approved": False}, headers=_h(tokens[actor]))
        assert r2.status_code == 200 and r2.json()["safety_team_approved"] is False


# ---------- Business approval still restricted ----------
class TestBusinessApprovalGating:
    def test_safety_cannot_business_approve(self, tokens, base_purchase):
        r = requests.patch(f"{API}/purchases/{base_purchase}/approval",
                           json={"approved": True}, headers=_h(tokens["safety_170"]))
        assert r.status_code == 403

    def test_editor_cannot_business_approve(self, tokens, base_purchase):
        r = requests.patch(f"{API}/purchases/{base_purchase}/approval",
                           json={"approved": True}, headers=_h(tokens["editor_16"]))
        assert r.status_code == 403

    def test_safety_cannot_bulk_approve(self, tokens, base_purchase):
        r = requests.patch(f"{API}/purchases/approval/bulk",
                           json={"ids": [base_purchase], "approved": True},
                           headers=_h(tokens["safety_170"]))
        assert r.status_code == 403

    def test_biz_can_business_approve(self, tokens, base_purchase):
        r = requests.patch(f"{API}/purchases/{base_purchase}/approval",
                           json={"approved": True}, headers=_h(tokens["biz_158"]))
        assert r.status_code == 200 and r.json()["business_manager_approved"] is True


# ---------- Safety user constraints on other endpoints ----------
class TestSafetyUserConstraints:
    def test_safety_delete_forbidden(self, tokens, base_purchase):
        r = requests.delete(f"{API}/purchases/{base_purchase}", headers=_h(tokens["safety_170"]))
        assert r.status_code == 403

    def test_safety_put_only_changes_safety_field(self, tokens, base_purchase):
        # Pre-set payment info by biz for baseline
        original = requests.get(f"{API}/purchases").json()
        original_item = next(i for i in original["items"] if i["id"] == base_purchase)
        # Safety attempts to change many fields
        put_payload = {
            "employee_id": "HACKED",
            "employee_name": "HACKED",
            "purchase_type": "Tech Device",
            "payment_mode": "Bank Transfer",
            "payment_by": "Mohammad Omer",
            "business_manager_approved": True,
            "safety_team_approved": True,
        }
        r = requests.put(f"{API}/purchases/{base_purchase}", json=put_payload, headers=_h(tokens["safety_170"]))
        assert r.status_code == 200, r.text
        # Verify via GET that only safety_team_approved changed
        listing = requests.get(f"{API}/purchases").json()
        item = next(i for i in listing["items"] if i["id"] == base_purchase)
        assert item["safety_team_approved"] is True
        assert item["employee_id"] == original_item["employee_id"]
        assert item["employee_name"] == original_item["employee_name"]
        assert item["purchase_type"] == original_item["purchase_type"]
        assert item["payment_mode"] == original_item["payment_mode"]
        assert item["payment_by"] == original_item["payment_by"]
        assert item["business_manager_approved"] == original_item["business_manager_approved"]


# ---------- Create-time approval gating ----------
class TestCreateApprovalGating:
    def _payload(self, eid, **over):
        p = {
            "employee_id": eid, "employee_name": "TEST C",
            "purchase_type": "Safety Shoes", "payment_mode": "Cash",
            "payment_by": "Jogy Joseph",
            "business_manager_approved": True, "safety_team_approved": True,
        }
        p.update(over)
        return p

    def test_safety_can_create_with_safety_true_only(self, tokens):
        r = requests.post(f"{API}/purchases", json=self._payload("TEST_C_SAFETY"), headers=_h(tokens["safety_170"]))
        assert r.status_code == 201
        j = r.json()
        assert j["safety_team_approved"] is True
        assert j["business_manager_approved"] is False  # not allowed for safety
        assert j["payment_mode"] == ""  # payment stripped
        assert j["payment_by"] == ""
        requests.delete(f"{API}/purchases/{j['id']}", headers=_h(tokens["biz_158"]))

    def test_editor_create_forces_both_false(self, tokens):
        r = requests.post(f"{API}/purchases", json=self._payload("TEST_C_ED"), headers=_h(tokens["editor_16"]))
        assert r.status_code == 201
        j = r.json()
        assert j["safety_team_approved"] is False
        assert j["business_manager_approved"] is False
        # editors CAN set payment
        assert j["payment_mode"] == "Cash"
        requests.delete(f"{API}/purchases/{j['id']}", headers=_h(tokens["biz_158"]))

    def test_anonymous_create_forces_false_and_no_payment(self):
        r = requests.post(f"{API}/purchases", json={
            "employee_id": "TEST_C_ANON", "employee_name": "Anon",
            "purchase_type": "Mobile Purchase", "payment_mode": "Cash",
            "payment_by": "Jogy Joseph",
            "business_manager_approved": True, "safety_team_approved": True,
        })
        assert r.status_code == 201
        j = r.json()
        assert j["safety_team_approved"] is False
        assert j["business_manager_approved"] is False
        assert j["payment_mode"] == "" and j["payment_by"] == ""

    def test_biz_can_create_with_both(self, tokens):
        r = requests.post(f"{API}/purchases", json=self._payload("TEST_C_BIZ"), headers=_h(tokens["biz_158"]))
        assert r.status_code == 201
        j = r.json()
        assert j["safety_team_approved"] is True and j["business_manager_approved"] is True
        requests.delete(f"{API}/purchases/{j['id']}", headers=_h(tokens["biz_158"]))


# ---------- Exports ----------
class TestExports:
    def test_xlsx_has_safety_column(self, tokens, base_purchase):
        # Approve safety on base_purchase
        requests.patch(f"{API}/purchases/{base_purchase}/safety-approval",
                       json={"approved": True}, headers=_h(tokens["safety_170"]))
        r = requests.get(f"{API}/purchases/export/xlsx")
        assert r.status_code == 200
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        ws = wb.active
        headers = [c.value for c in ws[1]]
        assert "Approved by Safety Team" in headers
        assert "Approved by Business Manager" in headers
        assert headers.index("Approved by Safety Team") == 6
        # Find our row and check value column
        safety_idx = headers.index("Approved by Safety Team")
        eid_idx = headers.index("Employee ID")
        found = False
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[eid_idx] == "TEST_S1":
                assert row[safety_idx] in ("Approved", "Not Approved")
                found = True
                break
        assert found, "TEST_S1 row not present in xlsx export"

    def test_pdf_200(self):
        r = requests.get(f"{API}/purchases/export/pdf")
        assert r.status_code == 200
        assert r.headers.get("content-type", "").startswith("application/pdf")
        assert len(r.content) > 500


# ---------- Regressions ----------
class TestRegressions:
    def test_anonymous_add_basics_only(self):
        r = requests.post(f"{API}/purchases", json={
            "employee_id": "TEST_ANON_R", "employee_name": "Anon R",
            "purchase_type": "Mobile Purchase",
        })
        assert r.status_code == 201
        j = r.json()
        assert j["payment_mode"] == "" and j["payment_by"] == ""
        assert j["business_manager_approved"] is False and j["safety_team_approved"] is False

    def test_editor_put_only_changes_payment(self, tokens, base_purchase):
        orig = next(i for i in requests.get(f"{API}/purchases").json()["items"] if i["id"] == base_purchase)
        r = requests.put(f"{API}/purchases/{base_purchase}", json={
            "employee_id": "IGN", "employee_name": "IGN",
            "purchase_type": "Tech Device",
            "payment_mode": "Credit Card", "payment_by": "Mohammad Omer",
            "business_manager_approved": True, "safety_team_approved": True,
        }, headers=_h(tokens["editor_16"]))
        assert r.status_code == 200
        item = next(i for i in requests.get(f"{API}/purchases").json()["items"] if i["id"] == base_purchase)
        assert item["payment_mode"] == "Credit Card"
        assert item["payment_by"] == "Mohammad Omer"
        assert item["employee_id"] == orig["employee_id"]
        assert item["purchase_type"] == orig["purchase_type"]
        # approvals unchanged by editor
        assert item["business_manager_approved"] == orig["business_manager_approved"]
        assert item["safety_team_approved"] == orig["safety_team_approved"]

    def test_no_token_put_delete_401(self, base_purchase):
        assert requests.put(f"{API}/purchases/{base_purchase}", json={
            "employee_id": "x", "employee_name": "x", "purchase_type": "Mobile Purchase"
        }).status_code == 401
        assert requests.delete(f"{API}/purchases/{base_purchase}").status_code == 401

    def test_pagination_and_search(self):
        r = requests.get(f"{API}/purchases", params={"page": 1, "page_size": 5})
        assert r.status_code == 200
        j = r.json()
        assert j["page_size"] == 5 and len(j["items"]) <= 5
        assert "total" in j

    def test_stats(self):
        r = requests.get(f"{API}/purchases/stats")
        assert r.status_code == 200
        for k in ("total", "approved", "pending", "mobile", "tech"):
            assert k in r.json()

    def test_biz_can_lift_both_approvals_via_put(self, tokens):
        # create fresh with both approvals
        cr = requests.post(f"{API}/purchases", json={
            "employee_id": "TEST_BOTH", "employee_name": "Both",
            "purchase_type": "Mobile Purchase", "payment_mode": "Cash",
            "payment_by": "Jogy Joseph",
            "business_manager_approved": True, "safety_team_approved": True,
        }, headers=_h(tokens["biz_158"]))
        pid = cr.json()["id"]
        try:
            # Lift both via PUT
            r = requests.put(f"{API}/purchases/{pid}", json={
                "employee_id": "TEST_BOTH", "employee_name": "Both",
                "purchase_type": "Mobile Purchase", "payment_mode": "Cash",
                "payment_by": "Jogy Joseph",
                "business_manager_approved": False, "safety_team_approved": False,
            }, headers=_h(tokens["biz_158"]))
            assert r.status_code == 200
            item = next(i for i in requests.get(f"{API}/purchases").json()["items"] if i["id"] == pid)
            assert item["business_manager_approved"] is False and item["safety_team_approved"] is False
        finally:
            requests.delete(f"{API}/purchases/{pid}", headers=_h(tokens["biz_158"]))

from fastapi import FastAPI, APIRouter, HTTPException, Query, UploadFile, File, Header, Depends
from fastapi.responses import StreamingResponse, Response
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import io
import csv
import logging
import requests
import jwt
from pathlib import Path
from pydantic import BaseModel, Field, ConfigDict, BeforeValidator
from typing import List, Optional, Annotated
from bson import ObjectId
import uuid
from datetime import datetime, timezone, timedelta

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

app = FastAPI()
api_router = APIRouter(prefix="/api")

ORG_NAME = "RLPC IT Assets Records"
PURCHASE_TYPES = ["Mobile Purchase", "Tech Device", "Safety Shoes"]
PAYMENT_MODES = ["Cash", "Credit Card", "Bank Transfer"]
PAYMENT_BY = ["Jogy Joseph", "Mohammad Omer", "Muhammad Khaleel", "Muhammad Abdullah"]

COLUMNS = ["Employee ID", "Employee Name", "Purchase Of", "Mode of Payment",
           "Payment By", "Approved by Business Manager", "Date"]

# ----- Auth (hardcoded single user) -----
JWT_SECRET = os.environ["JWT_SECRET"]
JWT_ALG = "HS256"
AUTH_USERNAME = os.environ["AUTH_USERNAME"]
AUTH_PASSWORD = os.environ["AUTH_PASSWORD"]


def _create_token(username: str) -> str:
    payload = {"sub": username, "exp": datetime.now(timezone.utc) + timedelta(days=7)}
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALG)


async def optional_user(authorization: Optional[str] = Header(None)) -> Optional[str]:
    if not authorization or not authorization.startswith("Bearer "):
        return None
    try:
        payload = jwt.decode(authorization[7:], JWT_SECRET, algorithms=[JWT_ALG])
        return payload.get("sub")
    except Exception:
        return None


# ----- Object storage (Emergent managed) -----
STORAGE_BASE = (os.environ.get("INTEGRATION_PROXY_URL") or "").strip() or "https://integrations.emergentagent.com"
STORAGE_URL = STORAGE_BASE.rstrip("/") + "/objstore/api/v1/storage"
EMERGENT_KEY = os.environ.get("EMERGENT_LLM_KEY")
APP_NAME = "rlpc-records"
storage_key = None


def init_storage(force: bool = False):
    global storage_key
    if storage_key and not force:
        return storage_key
    resp = requests.post(f"{STORAGE_URL}/init", json={"emergent_key": EMERGENT_KEY}, timeout=30)
    resp.raise_for_status()
    storage_key = resp.json()["storage_key"]
    return storage_key


def put_object(path: str, data: bytes, content_type: str) -> dict:
    key = init_storage()
    resp = requests.put(
        f"{STORAGE_URL}/objects/{path}",
        headers={"X-Storage-Key": key, "Content-Type": content_type},
        data=data, timeout=120,
    )
    if resp.status_code == 404:
        key = init_storage(force=True)
        resp = requests.put(
            f"{STORAGE_URL}/objects/{path}",
            headers={"X-Storage-Key": key, "Content-Type": content_type},
            data=data, timeout=120,
        )
    resp.raise_for_status()
    return resp.json()


def get_object(path: str):
    key = init_storage()
    resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    if resp.status_code == 404:
        key = init_storage(force=True)
        resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    resp.raise_for_status()
    return resp.content, resp.headers.get("Content-Type", "application/octet-stream")


def _validate_object_id(v):
    if isinstance(v, ObjectId):
        return str(v)
    return v


PyObjectId = Annotated[str, BeforeValidator(_validate_object_id)]


# ----- Models -----
class PurchaseBase(BaseModel):
    employee_id: str = Field(..., min_length=1)
    employee_name: str = Field(..., min_length=1)
    purchase_type: str
    payment_mode: str
    payment_by: str
    business_manager_approved: bool = False
    bill_path: Optional[str] = None
    bill_filename: Optional[str] = None


class PurchaseCreate(PurchaseBase):
    pass


class PurchaseUpdate(PurchaseBase):
    pass


class Purchase(PurchaseBase):
    model_config = ConfigDict(populate_by_name=True)
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    purchase_date: str
    created_at: str
    updated_at: str


def _validate_enums(p: PurchaseBase):
    if p.purchase_type not in PURCHASE_TYPES:
        raise HTTPException(status_code=422, detail="Invalid Purchase Of value")
    if p.payment_mode not in PAYMENT_MODES:
        raise HTTPException(status_code=422, detail="Invalid Mode of Payment value")
    if p.payment_by not in PAYMENT_BY:
        raise HTTPException(status_code=422, detail="Invalid Payment By value")


def _build_query(search, purchase_type, payment_mode, payment_by, approved, date_from, date_to):
    q = {}
    if search:
        q["$or"] = [
            {"employee_id": {"$regex": search, "$options": "i"}},
            {"employee_name": {"$regex": search, "$options": "i"}},
        ]
    if purchase_type:
        q["purchase_type"] = purchase_type
    if payment_mode:
        q["payment_mode"] = payment_mode
    if payment_by:
        q["payment_by"] = payment_by
    if approved is not None:
        q["business_manager_approved"] = approved
    date_q = {}
    if date_from:
        date_q["$gte"] = f"{date_from}T00:00:00"
    if date_to:
        date_q["$lte"] = f"{date_to}T23:59:59.999999"
    if date_q:
        q["purchase_date"] = date_q
    return q


async def _fetch_filtered(query):
    docs = await db.purchases.find(query, {"_id": 0}).sort("purchase_date", -1).to_list(length=100000)
    return docs


class LoginRequest(BaseModel):
    username: str
    password: str


# ----- Routes -----
@api_router.get("/")
async def root():
    return {"message": "RLPC IT Assets Records API", "org": ORG_NAME}


@api_router.post("/auth/login")
async def login(body: LoginRequest):
    if body.username != AUTH_USERNAME or body.password != AUTH_PASSWORD:
        raise HTTPException(status_code=401, detail="Invalid username or password")
    return {"token": _create_token(body.username), "username": body.username}


@api_router.get("/auth/me")
async def auth_me(user: Optional[str] = Depends(optional_user)):
    if not user:
        raise HTTPException(status_code=401, detail="Not authenticated")
    return {"username": user}


@api_router.get("/config")
async def get_config():
    return {
        "org_name": ORG_NAME,
        "purchase_types": PURCHASE_TYPES,
        "payment_modes": PAYMENT_MODES,
        "payment_by": PAYMENT_BY,
    }


@api_router.post("/purchases/upload-bill")
async def upload_bill(file: UploadFile = File(...)):
    ext = file.filename.split(".")[-1].lower() if "." in file.filename else "bin"
    path = f"{APP_NAME}/bills/{uuid.uuid4()}.{ext}"
    data = await file.read()
    result = put_object(path, data, file.content_type or "application/octet-stream")
    return {"bill_path": result["path"], "bill_filename": file.filename}


@api_router.get("/purchases/bill/{path:path}")
async def get_bill(path: str):
    data, content_type = get_object(path)
    return Response(content=data, media_type=content_type)


@api_router.post("/purchases", response_model=Purchase, status_code=201)
async def create_purchase(payload: PurchaseCreate, user: Optional[str] = Depends(optional_user)):
    _validate_enums(payload)
    if payload.business_manager_approved and not user:
        raise HTTPException(status_code=403, detail="Only an authorized user can mark a record as approved")
    now = datetime.now(timezone.utc).isoformat()
    obj = Purchase(**payload.model_dump(), purchase_date=now, created_at=now, updated_at=now)
    await db.purchases.insert_one(obj.model_dump())
    return obj


@api_router.get("/purchases")
async def list_purchases(
    search: Optional[str] = None,
    purchase_type: Optional[str] = None,
    payment_mode: Optional[str] = None,
    payment_by: Optional[str] = None,
    approved: Optional[bool] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=200),
):
    query = _build_query(search, purchase_type, payment_mode, payment_by, approved, date_from, date_to)
    total = await db.purchases.count_documents(query)
    skip = (page - 1) * page_size
    docs = await db.purchases.find(query, {"_id": 0}).sort("purchase_date", -1).skip(skip).limit(page_size).to_list(length=page_size)
    return {"items": docs, "total": total, "page": page, "page_size": page_size}


@api_router.get("/purchases/stats")
async def stats(
    search: Optional[str] = None,
    purchase_type: Optional[str] = None,
    payment_mode: Optional[str] = None,
    payment_by: Optional[str] = None,
    approved: Optional[bool] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
):
    query = _build_query(search, purchase_type, payment_mode, payment_by, approved, date_from, date_to)
    docs = await db.purchases.find(query, {"business_manager_approved": 1, "purchase_type": 1, "_id": 0}).to_list(length=100000)
    total = len(docs)
    approved_count = sum(1 for d in docs if d.get("business_manager_approved"))
    mobile = sum(1 for d in docs if d.get("purchase_type") == "Mobile Purchase")
    tech = sum(1 for d in docs if d.get("purchase_type") == "Tech Device")
    return {
        "total": total,
        "approved": approved_count,
        "pending": total - approved_count,
        "mobile": mobile,
        "tech": tech,
    }


@api_router.put("/purchases/{purchase_id}", response_model=Purchase)
async def update_purchase(purchase_id: str, payload: PurchaseUpdate, user: Optional[str] = Depends(optional_user)):
    _validate_enums(payload)
    existing = await db.purchases.find_one({"id": purchase_id}, {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail="Purchase record not found")
    if bool(payload.business_manager_approved) != bool(existing.get("business_manager_approved")) and not user:
        raise HTTPException(status_code=403, detail="Only an authorized user can change the approval status")
    update_doc = payload.model_dump()
    update_doc["updated_at"] = datetime.now(timezone.utc).isoformat()
    await db.purchases.update_one({"id": purchase_id}, {"$set": update_doc})
    merged = {**existing, **update_doc}
    return Purchase(**merged)


class ApprovalRequest(BaseModel):
    approved: bool


class BulkApprovalRequest(BaseModel):
    ids: List[str]
    approved: bool = True


@api_router.patch("/purchases/approval/bulk")
async def bulk_approval(body: BulkApprovalRequest, user: Optional[str] = Depends(optional_user)):
    if not user:
        raise HTTPException(status_code=403, detail="Only an authorized user can change the approval status")
    if not body.ids:
        return {"updated": 0}
    res = await db.purchases.update_many(
        {"id": {"$in": body.ids}},
        {"$set": {"business_manager_approved": body.approved, "updated_at": datetime.now(timezone.utc).isoformat()}},
    )
    return {"updated": res.modified_count}


@api_router.patch("/purchases/{purchase_id}/approval", response_model=Purchase)
async def set_approval(purchase_id: str, body: ApprovalRequest, user: Optional[str] = Depends(optional_user)):
    if not user:
        raise HTTPException(status_code=403, detail="Only an authorized user can change the approval status")
    existing = await db.purchases.find_one({"id": purchase_id}, {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail="Purchase record not found")
    update_doc = {
        "business_manager_approved": body.approved,
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.purchases.update_one({"id": purchase_id}, {"$set": update_doc})
    return Purchase(**{**existing, **update_doc})


@api_router.delete("/purchases/{purchase_id}")
async def delete_purchase(purchase_id: str):
    res = await db.purchases.delete_one({"id": purchase_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Purchase record not found")
    return {"success": True}


def _fmt_date(iso: str) -> str:
    try:
        dt = datetime.fromisoformat(iso)
        return dt.strftime("%Y-%m-%d %H:%M")
    except Exception:
        return iso or ""


def _row_values(d):
    return [
        d.get("employee_id", ""),
        d.get("employee_name", ""),
        d.get("purchase_type", ""),
        d.get("payment_mode", ""),
        d.get("payment_by", ""),
        "Approved" if d.get("business_manager_approved") else "Not Approved",
        _fmt_date(d.get("purchase_date", "")),
    ]


@api_router.get("/purchases/export/xlsx")
async def export_xlsx(
    search: Optional[str] = None,
    purchase_type: Optional[str] = None,
    payment_mode: Optional[str] = None,
    payment_by: Optional[str] = None,
    approved: Optional[bool] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
):
    query = _build_query(search, purchase_type, payment_mode, payment_by, approved, date_from, date_to)
    docs = await _fetch_filtered(query)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Purchase Records"

    header_fill = PatternFill(start_color="0A0A0A", end_color="0A0A0A", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(COLUMNS)
    for col_idx, _ in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=col_idx)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border

    for d in docs:
        ws.append(_row_values(d))

    for r in range(2, len(docs) + 2):
        for col_idx in range(1, len(COLUMNS) + 1):
            ws.cell(row=r, column=col_idx).border = border

    widths = [16, 24, 18, 18, 20, 26, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"purchase_records_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


@api_router.get("/purchases/export/pdf")
async def export_pdf(
    search: Optional[str] = None,
    purchase_type: Optional[str] = None,
    payment_mode: Optional[str] = None,
    payment_by: Optional[str] = None,
    approved: Optional[bool] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
):
    query = _build_query(search, purchase_type, payment_mode, payment_by, approved, date_from, date_to)
    docs = await _fetch_filtered(query)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        topMargin=15 * mm, bottomMargin=15 * mm, leftMargin=12 * mm, rightMargin=12 * mm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Title", parent=styles["Title"], fontSize=18, textColor=colors.HexColor("#0A0A0A"))
    sub_style = ParagraphStyle("Sub", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#6B7280"))
    cell_style = ParagraphStyle("Cell", parent=styles["Normal"], fontSize=8, leading=10)
    head_style = ParagraphStyle("Head", parent=styles["Normal"], fontSize=8, leading=10, textColor=colors.white, fontName="Helvetica-Bold")

    elements = [
        Paragraph(ORG_NAME, title_style),
        Spacer(1, 2 * mm),
        Paragraph(f"Exported: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}  |  Total records: {len(docs)}", sub_style),
        Spacer(1, 5 * mm),
    ]

    table_data = [[Paragraph(c, head_style) for c in COLUMNS]]
    for d in docs:
        table_data.append([Paragraph(str(v), cell_style) for v in _row_values(d)])

    col_widths = [30 * mm, 45 * mm, 32 * mm, 32 * mm, 40 * mm, 42 * mm, 36 * mm]
    tbl = Table(table_data, colWidths=col_widths, repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0A0A0A")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D0D0D0")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F7F8")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
    ]))
    elements.append(tbl)

    if not docs:
        elements.append(Paragraph("No records match the current filters.", sub_style))

    doc.build(elements)
    buf.seek(0)
    fname = f"purchase_records_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf"
    return StreamingResponse(
        buf, media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


HEADER_MAP = {
    "employee id": "employee_id",
    "employee_id": "employee_id",
    "employee name": "employee_name",
    "employee_name": "employee_name",
    "purchase of": "purchase_type",
    "purchase type": "purchase_type",
    "purchase_type": "purchase_type",
    "mode of payment": "payment_mode",
    "payment mode": "payment_mode",
    "payment_mode": "payment_mode",
    "payment by": "payment_by",
    "payment_by": "payment_by",
    "approved by business manager": "approved",
    "approved": "approved",
    "business_manager_approved": "approved",
    "date": "date",
    "purchase date": "date",
    "purchase_date": "date",
}

TRUE_VALUES = {"true", "yes", "1", "approved", "y", "t"}


def _parse_date(raw: str) -> Optional[str]:
    raw = (raw or "").strip()
    if not raw:
        return None
    try:
        return datetime.fromisoformat(raw).astimezone(timezone.utc).isoformat()
    except Exception:
        pass
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S"):
        try:
            dt = datetime.strptime(raw, fmt).replace(tzinfo=timezone.utc)
            return dt.isoformat()
        except Exception:
            continue
    return None


@api_router.post("/purchases/import")
async def import_purchases(file: UploadFile = File(...)):
    if not file.filename.lower().endswith(".csv"):
        raise HTTPException(status_code=422, detail="Please upload a .csv file")

    raw = await file.read()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")

    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise HTTPException(status_code=422, detail="CSV file is empty or has no header row")

    field_lookup = {}
    for h in reader.fieldnames:
        key = HEADER_MAP.get((h or "").strip().lower())
        if key:
            field_lookup[h] = key

    required = {"employee_id", "employee_name", "purchase_type", "payment_mode", "payment_by"}
    missing_headers = required - set(field_lookup.values())
    if missing_headers:
        pretty = {
            "employee_id": "Employee ID", "employee_name": "Employee Name",
            "purchase_type": "Purchase Of", "payment_mode": "Mode of Payment",
            "payment_by": "Payment By",
        }
        raise HTTPException(
            status_code=422,
            detail=f"Missing required columns: {', '.join(pretty[m] for m in missing_headers)}",
        )

    docs = []
    errors = []
    row_num = 1
    for row in reader:
        row_num += 1
        rec = {"business_manager_approved": False}
        purchase_date = None
        for header, key in field_lookup.items():
            val = (row.get(header) or "").strip()
            if key == "approved":
                rec["business_manager_approved"] = val.lower() in TRUE_VALUES
            elif key == "date":
                if val:
                    purchase_date = _parse_date(val)
                    if purchase_date is None:
                        errors.append(f"Row {row_num}: could not parse date '{val}'")
            else:
                rec[key] = val

        row_errors = []
        for f in ("employee_id", "employee_name"):
            if not rec.get(f):
                row_errors.append(f.replace("_", " ").title())
        if rec.get("purchase_type") not in PURCHASE_TYPES:
            row_errors.append(f"invalid Purchase Of '{rec.get('purchase_type', '')}'")
        if rec.get("payment_mode") not in PAYMENT_MODES:
            row_errors.append(f"invalid Mode of Payment '{rec.get('payment_mode', '')}'")
        if rec.get("payment_by") not in PAYMENT_BY:
            row_errors.append(f"invalid Payment By '{rec.get('payment_by', '')}'")

        if row_errors:
            errors.append(f"Row {row_num}: {', '.join(row_errors)}")
            continue

        now = datetime.now(timezone.utc).isoformat()
        obj = Purchase(
            employee_id=rec["employee_id"], employee_name=rec["employee_name"],
            purchase_type=rec["purchase_type"], payment_mode=rec["payment_mode"],
            payment_by=rec["payment_by"], business_manager_approved=rec["business_manager_approved"],
            purchase_date=purchase_date or now, created_at=now, updated_at=now,
        )
        docs.append(obj.model_dump())

    if docs:
        await db.purchases.insert_many(docs)

    return {"imported": len(docs), "failed": len(errors), "errors": errors[:50]}


app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)


@app.on_event("startup")
async def _startup():
    try:
        init_storage()
        logger.info("Object storage initialized")
    except Exception as e:
        logger.error(f"Storage init failed: {e}")


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
